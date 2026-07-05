"""Format-only parsing for the universal ingest adapter (Bounty 02).

Deliberately dumb: we turn an uploaded file into a raw grid of cells (or a JSON
value) and hand THAT to the LLM. Interpreting the messy shape — which row is the
header, which column is the site, what the file even *is* — is the model's job,
not a per-source connector's. That "no bespoke parser per source" property is the
"universal translator" the bounty asks for.

`parse_upload` never raises on messy content; it raises only on a genuinely
unreadable/unsupported file, so the router can return a clean 400.
"""

import csv
import io
import json

# Keep prompts bounded — a stock/revenue export is tens of rows, not thousands.
MAX_PREVIEW_ROWS = 40
MAX_PREVIEW_CHARS = 8000


class UnsupportedFileError(ValueError):
    pass


def _ext(filename: str) -> str:
    return filename.rsplit(".", 1)[-1].lower() if "." in filename else ""


def _grid_to_text(rows: list[list[str]]) -> str:
    """Render a cell grid as CSV-ish text for the prompt (row-numbered so the
    model can point at 'the header is row 2' etc.)."""
    lines = []
    for i, row in enumerate(rows[:MAX_PREVIEW_ROWS]):
        cells = ", ".join("" if c is None else str(c) for c in row)
        lines.append(f"{i}: {cells}")
    text = "\n".join(lines)
    if len(rows) > MAX_PREVIEW_ROWS:
        text += f"\n… ({len(rows) - MAX_PREVIEW_ROWS} more rows omitted)"
    return text[:MAX_PREVIEW_CHARS]


def _parse_csv(data: bytes) -> dict:
    text = data.decode("utf-8-sig", errors="replace")
    rows = list(csv.reader(io.StringIO(text)))
    return {"kind": "csv", "rows": rows, "grid": _grid_to_text(rows)}


def _parse_json(data: bytes) -> dict:
    text = data.decode("utf-8-sig", errors="replace")
    try:
        value = json.loads(text)
    except json.JSONDecodeError as exc:
        raise UnsupportedFileError(f"file is not valid JSON: {exc}") from exc
    pretty = json.dumps(value, indent=2, default=str)
    return {"kind": "json", "rows": value, "grid": pretty[:MAX_PREVIEW_CHARS]}


def _parse_xlsx(data: bytes) -> dict:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises assorted errors on bad files
        raise UnsupportedFileError(f"could not read xlsx: {exc}") from exc
    ws = wb.active
    rows = [
        ["" if cell is None else cell for cell in row]
        for row in ws.iter_rows(values_only=True)
    ]
    wb.close()
    return {"kind": "xlsx", "rows": rows, "grid": _grid_to_text(rows)}


def parse_upload(filename: str, data: bytes) -> dict:
    """Return {filename, kind, rows, grid}. `grid` is a bounded text preview the
    LLM reads; `rows` is the structured content (list-of-lists for tabular, the
    parsed value for JSON). Raises UnsupportedFileError on unreadable input."""
    if not data:
        raise UnsupportedFileError("empty file")
    ext = _ext(filename)
    if ext == "csv":
        parsed = _parse_csv(data)
    elif ext == "json":
        parsed = _parse_json(data)
    elif ext in ("xlsx", "xlsm"):
        parsed = _parse_xlsx(data)
    else:
        raise UnsupportedFileError(f"unsupported file type {ext!r}; use CSV, JSON, or xlsx")
    return {"filename": filename, **parsed}
